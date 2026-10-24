"""Модульные документы: находим поля для заполнения по сноскам и собираем черновик.

Поле (слот) — прочерк «____» или вариант в квадратных скобках «[а / б]».
Сноска ^[N] относится ко всем полям абзаца между предыдущей сноской и ней.
Пояснения к сноскам берутся из раздела «Примечания по заполнению».
После заполнения чистим документ по п. 3.3 Методических рекомендаций.
"""
from __future__ import annotations

import copy
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path

import docx
from docx.document import Document
from docx.text.paragraph import Paragraph

from . import config

BLANK_RE = re.compile(r"«_+»\s*_+\s*20_+\s*г\.|«_+»\s*_+|_{3,}|\[[^\[\]]*?/[^\[\]]*?\]")
MARK_RE = re.compile(r"^\s*\[(\d+)\]\s*$")
DATE_RE = re.compile(r"^«_+»")
YEAR_RE = re.compile(r"20_+\s*г\.$")
NOTE_RE = re.compile(r"^(\d+)\.\s+(.+?)\.\s+(.*)$")
MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
          "сентября", "октября", "ноября", "декабря"]

# Какой строке Матрицы визирования соответствует модуль (группа, строка)
MATRIX_ROW = {
    "Модуль_приказа_о_вводе_ИС": ("Внутренние документы", 3),
    "Модуль_приказа_о_приеме_и_передаче_дел": ("Внутренние документы", 2),
    "Модуль_приказа_о_назначении_ответственных": ("Внутренние документы", 4),
    "Модуль_приказа_об_утверждении_ЛНА": ("Внутренние документы", 4),
    "Модуль_распоряжения_о_внутреннем_аудите": ("Внутренние документы", 7),
    "Модуль_распоряжения_о_проектном_комитете": ("Внутренние документы", 6),
    "Модуль_распоряжения_о_рабочей_группе": ("Внутренние документы", 6),
}


@dataclass
class Slot:
    id: str          # «s3»
    kind: str        # text | date | choice | placeholder
    footnote: int | None
    label: str       # короткое название из примечаний
    hint: str        # пояснение из примечаний
    context: str     # текст перед полем, чтобы было понятно, куда оно встанет
    options: list[str] | None = None


def _all_paragraphs(d: Document):
    """Абзацы тела и ячеек таблиц в порядке документа."""
    body = d.element.body
    for el in body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            yield Paragraph(el, d)
        elif tag == "tbl":
            for p_el in el.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                yield Paragraph(p_el, d)


def _is_mark(run) -> int | None:
    m = MARK_RE.match(run.text)
    if m and (run.font.superscript or run.text.strip().startswith("[")):
        return int(m.group(1))
    return None


def _notes(d: Document) -> dict[int, tuple[str, str]]:
    notes, started = {}, False
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("Примечания по заполнению"):
            started = True
            continue
        if started:
            m = NOTE_RE.match(t)
            if m:
                notes[int(m.group(1))] = (m.group(2), m.group(3))
    return notes


def _body_paragraphs(d: Document):
    """Только сам документ: без грифа «УТВЕРЖДАЮ» и без примечаний."""
    for p in _all_paragraphs(d):
        if p.text.strip().startswith("Примечания по заполнению"):
            return
        if _in_approval_table(p):
            continue
        yield p


def _in_approval_table(p: Paragraph) -> bool:
    tbl = p._p
    while tbl is not None and not tbl.tag.endswith("}tbl"):
        tbl = tbl.getparent()
    if tbl is None:
        return False
    texts = "".join(t.text or "" for t in tbl.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
    return "УТВЕРЖДАЮ" in texts


def _fmt(r):
    c = r.font.color.rgb if r.font.color is not None and r.font.color.type else None
    return (r.bold, r.italic, r.underline, r.font.superscript, c, r.font.size, r.font.name)


def _merge_runs(p: Paragraph):
    """Склеиваем соседние run с одинаковым форматированием: Word режет ««___» ___ 20__ г.» на части."""
    runs = p.runs
    i = 0
    while i < len(runs) - 1:
        a, b = runs[i], runs[i + 1]
        if _is_mark(a) is None and _is_mark(b) is None and _fmt(a) == _fmt(b):
            a.text += b.text
            b._r.getparent().remove(b._r)
            runs = p.runs
        else:
            i += 1


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _el_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(W + "t")).strip()


def _row_label(p: Paragraph) -> str:
    """Для ячейки таблицы — текст первой ячейки строки («Руководитель»),
    а если она сама поле — заголовок столбца и номер строки."""
    tc = p._p
    while tc is not None and not tc.tag.endswith("}tc"):
        tc = tc.getparent()
    if tc is None:
        return ""
    tr = tc.getparent()
    cells = list(tr.iterchildren(W + "tc"))
    first = _el_text(cells[0])
    if first and not re.fullmatch(r"[_\s]+", first) and cells[0] is not tc:
        return first
    tbl = tr.getparent()
    rows = list(tbl.iterchildren(W + "tr"))
    head = list(rows[0].iterchildren(W + "tc"))
    col = cells.index(tc)
    if col < len(head) and rows[0] is not tr:
        return f"{_el_text(head[col])} (строка {rows.index(tr)})"
    return ""


def _scan(d: Document):
    """Возвращает список (абзац, [(run_idx, match|None, footnote)]) для всех полей."""
    out = []
    for p in _body_paragraphs(d):
        _merge_runs(p)
        runs = p.runs
        tokens = []  # ("blank", ri, m) | ("mark", ri, n) в порядке текста
        for ri, r in enumerate(runs):
            n = _is_mark(r)
            if n is not None:
                tokens.append(("mark", ri, n))
            else:
                tokens += [("blank", ri, m) for m in BLANK_RE.finditer(r.text)]
        items, pending = [], []
        for ti, (kind, ri, x) in enumerate(tokens):
            if kind == "used":
                continue
            if kind == "blank":
                pending.append((ri, x))
                continue
            if pending:
                items += [(b_ri, b_m, x) for b_ri, b_m in pending]
                pending = []
                continue
            # сноска без прочерков перед ней: либо поле сразу после неё
            # («с [8] «__» ___ 20__ г.»), либо текст-образец перед ней («И. О. Фамилия [8]»)
            rest = tokens[ti + 1:]
            if rest and rest[0][0] == "blank" and not any(t[0] == "mark" for t in rest):
                items.append((rest[0][1], rest[0][2], x))
                tokens[ti + 1] = ("used",) + rest[0][1:]
                continue
            prev = ri - 1
            while prev >= 0 and not runs[prev].text.strip():
                prev -= 1
            if prev >= 0:
                items.append((prev, None, x))
        items += [(b_ri, b_m, None) for b_ri, b_m in pending]
        if items:
            items.sort(key=lambda it: (it[0], it[1].start() if it[1] else -1))
            out.append((p, items))
    return out


def _signature_line(p: Paragraph) -> bool:
    """Линия для подписи в блоке подписанта: её не заполняем, а удаляем."""
    if not re.fullmatch(r"[_\s]+", p.text):
        return False
    el = p._p
    while el is not None and not el.tag.endswith("}tbl"):
        el = el.getparent()
    if el is None:
        return False
    texts = "".join(t.text or "" for t in el.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
    return "ВИЗУАЛИЗАЦИЯ ЭП" in texts or "Фамилия" in texts


def list_modules() -> list[dict]:
    res = []
    for path in sorted(config.MODULES_DIR.glob("*.docx")):
        d = docx.Document(str(path))
        heading = next((p.text.strip() for p in d.paragraphs
                        if re.match(r"^О[б]?\s", p.text.strip())), "")
        kind = "Приказ" if "приказа" in path.stem else "Распоряжение"
        res.append({"id": path.stem, "title": path.stem.replace("Модуль_", "").replace("_", " "),
                    "kind": kind, "heading": re.sub(r"\s*_{3,}.*|\s*\[\d+\]", " …", heading).strip()})
    return res


def module_path(module_id: str) -> Path:
    p = config.MODULES_DIR / f"{module_id}.docx"
    if not p.exists() or p.parent != config.MODULES_DIR:
        raise KeyError(module_id)
    return p


def get_slots(module_id: str) -> list[Slot]:
    d = docx.Document(str(module_path(module_id)))
    notes = _notes(d)
    slots, k = [], 0
    for p, items in _scan(d):
        if _signature_line(p):
            continue
        for ri, m, fn in items:
            k += 1
            run_text = p.runs[ri].text
            label, hint = notes.get(fn, ("", "")) if fn else ("", "")
            before = "".join(r.text for r in p.runs[:ri]) + (run_text[: m.start()] if m else "")
            before = re.sub(r"\s*\[\d+\]", "", BLANK_RE.sub("…", before)).strip()
            row = _row_label(p)
            ctx = (f"{row}: " if row else "") + before[-70:]
            if m is None:
                kind, opts = "placeholder", None
                ctx = (ctx + " " + run_text.strip()).strip()[-70:]
            elif DATE_RE.match(m.group()):
                kind, opts = "date", None
            elif m.group().startswith("["):
                kind = "choice"
                opts = [o.strip() for o in m.group()[1:-1].split("/")]
            else:
                kind, opts = "text", None
            if not label and row:
                label = row
            if not label:
                label = "Заголовок" if re.match(r"^О[б]?\s", p.text.strip()) else "Поле"
            slots.append(Slot(f"s{k}", kind, fn, label, hint, ctx, opts))
    return slots


def format_date(value: str, with_year: bool = True) -> str:
    v = value.strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            dt = datetime.strptime(v, fmt).date()
            if not with_year:
                return f"«{dt.day:02d}» {MONTHS[dt.month - 1]}"
            return f"«{dt.day:02d}» {MONTHS[dt.month - 1]} {dt.year} г."
        except ValueError:
            pass
    return v


def fill(module_id: str, values: dict[str, str]) -> tuple[Path, list[str]]:
    """Заполняет копию модуля и чистит её. Возвращает путь к файлу и предупреждения."""
    src = module_path(module_id)
    d = docx.Document(str(src))
    slots = {s.id: s for s in get_slots(module_id)}
    warnings, k = [], 0

    for p, items in _scan(d):
        if _signature_line(p):
            for r in p.runs:
                r.text = ""
            continue
        # заменяем справа налево, чтобы не сбить позиции в одном run
        repl = []
        for ri, m, fn in items:
            k += 1
            sid = f"s{k}"
            val = (values.get(sid) or "").strip()
            s = slots.get(sid)
            if s and s.kind == "date" and val and m is not None:
                formatted = format_date(val, with_year=bool(YEAR_RE.search(m.group())))
                if formatted == val:
                    warnings.append(f"Дата «{val}» в поле «{s.label}» не распознана — укажите в формате ДД.ММ.ГГГГ")
                val = formatted
            if not val:
                warnings.append(f"Не заполнено поле «{s.label if s else sid}»: {s.context if s else ''}")
                val = m.group() if m else p.runs[ri].text
            repl.append((ri, m, val))
        for ri, m, val in sorted(repl, key=lambda x: (x[0], x[1].start() if x[1] else 0), reverse=True):
            r = p.runs[ri]
            r.text = r.text[: m.start()] + val + r.text[m.end():] if m else val
        for r in p.runs:
            if _is_mark(r) is not None:
                r.text = ""

    _cleanup(d)
    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = config.OUTPUT_DIR / f"{module_id.replace('Модуль_', 'Проект_')}_{datetime.now():%Y%m%d_%H%M%S}.docx"
    d.save(str(out))
    return out, warnings


def _cleanup(d: Document):
    """П. 3.3 Методических рекомендаций: убрать гриф, пометку модуля, примечания."""
    body = d.element.body
    for tbl in list(body.iterchildren("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")):
        texts = "".join(t.text or "" for t in tbl.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
        if "УТВЕРЖДАЮ" in texts:
            body.remove(tbl)
    remove = False
    for p in list(d.paragraphs):
        t = p.text.strip()
        if t.startswith("Примечания по заполнению"):
            remove = True
        if remove or t == "МОДУЛЬНЫЙ ДОКУМЕНТ":
            p._p.getparent().remove(p._p)


# ---- Проверка загруженного проекта ------------------------------------------------

def check_draft(path: Path) -> list[dict]:
    """Простые проверки проекта по п. 3.3 Методических рекомендаций и гл. 8 Инструкции."""
    d = docx.Document(str(path))
    text = "\n".join(p.text for p in _all_paragraphs(d))
    issues = []

    def add(ok, msg, ref):
        if not ok:
            issues.append({"message": msg, "ref": ref})

    add(not re.search(r"_{3,}", text), "В документе остались прочерки (незаполненные поля).",
        "Методические рекомендации, п. 3.3")
    add(not re.search(r"\[\d+\]", text), "Остались номера примечаний вида [1].",
        "Методические рекомендации, п. 3.3")
    add(not re.search(r"\[[^\]]*/[^\]]*\]", text), "Остались неиспользованные варианты в квадратных скобках.",
        "Методические рекомендации, п. 3.3")
    add("МОДУЛЬНЫЙ ДОКУМЕНТ" not in text, "Не удалена пометка «МОДУЛЬНЫЙ ДОКУМЕНТ».",
        "Методические рекомендации, п. 3.3")
    add("УТВЕРЖДАЮ" not in text, "Не удалён гриф утверждения «УТВЕРЖДАЮ».",
        "Методические рекомендации, п. 3.3")
    add("Примечания по заполнению" not in text, "Не удалены примечания по заполнению.",
        "Методические рекомендации, п. 3.3")
    add("ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР" in text, "Удалена служебная зона для даты и регистрационного номера.",
        "Методические рекомендации, п. 3.3")
    add("ВИЗУАЛИЗАЦИЯ ЭП" in text, "Удалена зона визуализации электронной подписи.",
        "Методические рекомендации, п. 3.3")
    add(re.search(r"^О[б]?\s", text, re.M) is not None, "Не найден заголовок к тексту (начинается с «О…»).",
        "п. 8 Инструкции (реквизит «заголовок к тексту»)")
    return issues
