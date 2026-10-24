"""Проект приказа о вводе ИС с типичными ошибками — для показа проверки на защите.

Запуск: python -m tools.make_bad_draft [путь.docx]
"""
import copy
import sys
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

from app import modules

VALUES = {
    "s1": "системы «Электронный архив документов»", "s2": "протокола приёмочных испытаний от 10.10.2026 № 15",
    "s3": "«Электронный архив документов»", "s4": "«ЭАД", "s6": "технического задания от 01.03.2026", "s7": "ЭАД",
    "s9": "Департаменту ДОУ", "s10": "ЭАД", "s11": "хранения и поиска электронных документов", "s12": "ЭАД",
    "s13": "ЭАД", "s14": "директор Департамента информационных технологий Петров П. П.",
}  # s5 (разработанного / измененного) и s8 (дата ввода) намеренно пустые


def make(out: Path) -> Path:
    src, _ = modules.fill("Модуль_приказа_о_вводе_ИС", VALUES)
    d = docx.Document(str(src))
    for p in d.paragraphs:
        if p.text.startswith("ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР"):
            p._p.getparent().remove(p._p)  # удалена служебная зона
        elif "обеспечить применение" in p.text:
            for r in p.runs:  # изменена постоянная формулировка модуля
                r.text = r.text.replace("обеспечить применение", "организовать использование")
        elif p.text.startswith("3.1."):
            p.runs[-1].text += " [7]"  # осталась сноска
    first = d.paragraphs[0]
    mark = copy.deepcopy(first._p)
    first._p.addprevious(mark)
    mp = Paragraph(mark, first._parent)
    for r in mp.runs[1:]:
        r.text = ""
    mp.runs[0].text = "МОДУЛЬНЫЙ ДОКУМЕНТ"  # не удалена пометка модуля
    d.save(str(out))
    return out


if __name__ == "__main__":
    out = make(Path(sys.argv[1] if len(sys.argv) > 1 else "Проект_с_ошибками_ввод_ИС.docx"))
    for i in modules.analyze_draft(out)["issues"]:
        print("-", i["message"], f"[{i['ref']}]")
